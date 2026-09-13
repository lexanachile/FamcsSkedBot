"""Check lesson classification on Excel cells without Google credentials."""
import ast
import re
import unittest
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet


source = ast.parse((Path(__file__).parents[1] / 'parser.py').read_text(encoding='utf-8'))
names = {'is_dp_title', 'extract_class', 'get_rect_values', 'get_cell_value',
         'get_merged_range_for_cell', 'parse_time_range'}
namespace = dict(re=re, Any=Any, Dict=Dict, List=List, Optional=Optional,
                 Tuple=Tuple, Worksheet=Worksheet, STANDARD_SLOT_HEIGHT=4)
exec(compile(ast.Module(body=[node for node in source.body
                             if isinstance(node, ast.FunctionDef) and node.name in names],
                        type_ignores=[]), 'parser.py', 'exec'), namespace)


class DpTests(unittest.TestCase):
    def extract(self, ws, **kwargs):
        return namespace['extract_class'](
            ws, (1, 1, 4, 2), '1', 3, 1, '08:15-09:35', **kwargs)[0]

    def test_common_dp_spelling_variants(self):
        for title in ('ДП-1', 'ДП-2', 'дп1', 'Дп-2', ' дП - 12 ', 'ДП–2', 'ДП‑1'):
            with self.subTest(title=title):
                ws = Workbook().active
                ws.merge_cells('A1:B1')
                ws['A1'] = title
                lesson = self.extract(ws)
                self.assertEqual((lesson['isCommon'], lesson['isLecture']), (1, 1))
                self.assertEqual(lesson['classTitleA'], title.strip())

    def test_separate_subgroups_are_practice(self):
        for titles in (('ДП-1', None), (None, 'дп2'), ('ДП-1', 'ДП-2'), ('ДП-1', 'ДП-1')):
            with self.subTest(titles=titles):
                ws = Workbook().active
                for col, title in enumerate(titles, 1):
                    if title:
                        ws.cell(1, col, title)
                        ws.cell(3, col, f'Преподаватель {col}')
                        ws.cell(4, col, str(100 + col))
                lesson = self.extract(ws)
                self.assertEqual((lesson['isCommon'], lesson['isLecture']), (0, 0))

    def test_stretched_title_with_one_subgroup_is_practice(self):
        for col in (1, 2):
            with self.subTest(col=col):
                ws = Workbook().active
                ws.merge_cells('A1:B1')
                ws['A1'] = 'ДП-1'
                ws.cell(3, col, 'Иванов')
                ws.cell(4, col, '101')
                lesson = self.extract(ws)
                self.assertEqual((lesson['isCommon'], lesson['isLecture']), (0, 0))

    def test_merged_layouts_and_single_group_flow(self):
        for merged in ('A1:B1', 'A1:B2', 'A1:B3', 'A1:B4'):
            with self.subTest(merged=merged):
                ws = Workbook().active
                ws.merge_cells(merged)
                ws['A1'] = 'ДП-2'
                lesson = self.extract(ws, lecture_coverage_start=1,
                                      lecture_coverage_end=2, is_single_group_flow=True)
                self.assertEqual((lesson['isCommon'], lesson['isLecture']), (1, 1))

    def test_other_common_subjects_stay_practice(self):
        for title in ('Физическая культура', 'Алгебра', 'ДП-1 дополнительное', 'АДП-1'):
            for merged in ('A1:B1', 'A1:B4'):
                with self.subTest(title=title, merged=merged):
                    ws = Workbook().active
                    ws.merge_cells(merged)
                    ws['A1'] = title
                    self.assertEqual(self.extract(ws)['isLecture'], 0)

    def test_regular_flow_lecture_keeps_metadata(self):
        ws = Workbook().active
        ws.merge_cells('A1:D2')
        ws['A1'] = 'Алгебра'
        ws['A3'] = 'Иванов'
        ws['C4'] = '130 ФМО'
        lesson = self.extract(ws, lecture_coverage_start=1, lecture_coverage_end=4)
        self.assertEqual(lesson['isLecture'], 1)
        self.assertEqual(lesson['professorNameA'], 'Иванов')
        self.assertEqual(lesson['classroomA'], '130 ФМО')


if __name__ == '__main__':
    unittest.main()
