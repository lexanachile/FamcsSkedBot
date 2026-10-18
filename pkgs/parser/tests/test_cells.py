"""Check module startup and Excel values without network requests or real secrets."""
import io
import os
import runpy
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook


class CellTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Execute module definitions too: AST-only tests miss broken imports.
        with patch.dict(os.environ, {
            'GOOGLE_SERVICE_ACCOUNT_JSON': '{}',
            'DRIVE_FOLDER_ID': 'test-folder',
            'BACKEND_URL': '',
            'BACKEND_AUTH_TOKEN': '',
        }), redirect_stdout(io.StringIO()):
            cls.namespace = runpy.run_path(
                str(Path(__file__).parents[1] / 'parser.py'))

    def test_cell_values(self):
        ws = Workbook().active
        for value, expected in (
            (None, None),
            (101, '101'),
            (101.0, '101'),
            (0.0, '0'),
            (-101.0, '-101'),
            (101.5, '101.5'),
            ('0101', '0101'),
            ('130 ФМО', '130 ФМО'),
            (True, 'True'),
        ):
            with self.subTest(value=value):
                ws['A1'] = value
                self.assertEqual(
                    self.namespace['get_cell_value'](ws, 1, 1), expected)

    def test_merged_numeric_cell(self):
        ws = Workbook().active
        ws.merge_cells('A1:B2')
        ws['A1'] = 101.0
        self.assertEqual(self.namespace['get_cell_value'](ws, 2, 2), '101')

    def test_empty_merged_cell(self):
        ws = Workbook().active
        ws.merge_cells('A1:B2')
        self.assertIsNone(self.namespace['get_cell_value'](ws, 2, 2))


if __name__ == '__main__':
    unittest.main()
